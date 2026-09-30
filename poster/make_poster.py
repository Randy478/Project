"""Erzeugt das A3-Poster "Die geneigte Ebene" als .pptx (fuer den Import in Canva)."""
import math

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Mm, Pt

BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xD0, 0x00, 0x00)
RAMP = RGBColor(0xDD, 0xDD, 0xDD)
FONT = "Arial"

prs = Presentation()
prs.slide_width, prs.slide_height = Mm(297), Mm(420)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def text(x, y, w, h, paras, size=18, color=BLACK, bold=False, align=PP_ALIGN.LEFT):
    """paras: Liste von Strings oder (String, dict) mit Abweichungen."""
    box = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = Mm(1)
    for i, p in enumerate(paras):
        s, opt = (p, {}) if isinstance(p, str) else p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = opt.get("align", align)
        para.space_after = Pt(opt.get("after", 4))
        run = para.add_run()
        run.text = s
        run.font.name = FONT
        run.font.size = Pt(opt.get("size", size))
        run.font.bold = opt.get("bold", bold)
        run.font.color.rgb = opt.get("color", color)
    return box


def rect(x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, dash=False, lw=1.5):
    s = slide.shapes.add_shape(shape, Mm(x), Mm(y), Mm(w), Mm(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    style = s._element.find(qn("p:style"))
    if style is not None:  # Theme-Schatten entfernen
        s._element.remove(style)
    return s


def arrow(x1, y1, x2, y2, color, width=3):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Mm(x1), Mm(y1), Mm(x2), Mm(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    return c


def heading(x, y, w, s):
    text(x, y, w, 12, [s], size=28, bold=True)


def cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        ln = etree.SubElement(tcPr, qn(side), w="12700")
        fill = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr"), val="000000")


M = 15          # Rand
W = 297 - 2 * M  # nutzbare Breite
COL = (W - 10) / 2

# ---------- Kopf ----------
text(M, 12, W, 25, ["Die geneigte Ebene"], size=60, bold=True, align=PP_ALIGN.CENTER)
text(M, 36, W, 12, ["und die Goldene Regel der Mechanik"], size=30, align=PP_ALIGN.CENTER)
text(M, 50, W, 8, ["von Can, Antonio, Florian und Philly (MCG)"], size=16,
     align=PP_ALIGN.CENTER)

# ---------- Was ist eine geneigte Ebene? ----------
y = 68
heading(M, y, COL, "Was ist das?")
text(M, y + 14, COL, 60, [
    "• eine schräge Fläche, z. B. eine Rampe",
    "• man zieht einen Körper die Rampe hoch, statt ihn senkrecht hochzuheben",
    "• dafür braucht man weniger Kraft",
    "• sie ist eine kraftumformende Einrichtung",
])

# ---------- Goldene Regel ----------
x2 = M + COL + 10
heading(x2, y, COL, "Die Goldene Regel")
text(x2, y + 14, COL, 60, [
    ("„Was man an Kraft spart, muss man mit Strecke bezahlen.“", {"bold": True}),
    "• je länger die Strecke, desto kleiner die Kraft",
    "• die Arbeit W = F · s bleibt immer gleich",
    "• Energie kann man so nicht sparen, nur Kraft",
])

# ---------- Skizze Rampe ----------
y = 146
heading(M, y, W, "Beispiel an der Rampe")
tx, ty, tw, th = M + 15, y + 22, 120, 45        # Dreieck (Rampe)
ramp = rect(tx, ty, tw, th, fill=RAMP, line=BLACK, shape=MSO_SHAPE.RIGHT_TRIANGLE)
# Dreieck spiegeln, damit die Rampe nach rechts ansteigt
ramp._element.spPr.find(qn("a:xfrm")).set("flipH", "1")

ang = math.degrees(math.atan(th / tw))
# Wagen auf der Rampe (bei 45 % der Strecke)
t = 0.45
cx, cy = tx + t * tw, ty + th - t * th
cw, ch = 18, 8
off = ch / 2 + 0.5
px = cx - math.sin(math.radians(ang)) * off
py = cy - math.cos(math.radians(ang)) * off
car = rect(px - cw / 2, py - ch / 2, cw, ch, fill=WHITE, line=BLACK)
car.rotation = -ang
# Kraftpfeil entlang der Rampe
d = 32
arrow(px + math.cos(math.radians(ang)) * 10, py - math.sin(math.radians(ang)) * 10,
      px + math.cos(math.radians(ang)) * d, py - math.sin(math.radians(ang)) * d, RED)
text(px + 14, py - 22, 30, 8, ["F"], size=20, bold=True, color=RED)
# Strecke s
s_lbl = text(tx + tw * 0.55, ty + th * 0.55, 20, 8, ["s"], size=20, bold=True)
s_lbl.rotation = -ang
# Hoehe h
arrow(tx + tw + 6, ty + th, tx + tw + 6, ty, BLACK, width=1.5)
text(tx + tw + 8, ty + th / 2 - 4, 30, 8, ["h = 1 m"], size=18)

text(M + 175, y + 16, W - 175, 60, [
    "Ein Wagen (1 kg) wird 1 m hoch gebracht.",
    ("Senkrecht hochheben:", {"bold": True, "after": 0}),
    "W = 9,8 N · 1,0 m = 9,8 J",
    ("Über die Rampe:", {"bold": True, "after": 0}),
    "W = 4,1 N · 2,4 m = 9,8 J",
    "Beide Male ist die Arbeit gleich.",
])

# ---------- Experiment ----------
y = 224
heading(M, y, W, "Experiment")
text(M, y + 14, 125, 80, [
    ("Material:", {"bold": True, "after": 0}),
    "Wagen (1 kg), Rampe, Kraftmesser",
    ("Aufbau:", {"bold": True, "after": 0}),
    "Die Rampe endet immer in 1 m Höhe.",
    ("Durchführung:", {"bold": True, "after": 0}),
    "Wir ziehen den Wagen langsam die Rampe hoch und lesen die Kraft F ab. "
    "Die Strecke s messen wir. Das machen wir mit verschieden steilen Rampen "
    "und rechnen jedes Mal F · s aus.",
])

# Messwerte-Tabelle
rows = [("s in m", "F in N", "F · s in N·m"),
        ("2,4", "4,1", "9,8"), ("2,0", "4,9", "9,8"), ("1,6", "6,2", "9,9"),
        ("1,2", "8,1", "9,7"), ("1,0", "9,9", "9,9")]
gf = slide.shapes.add_table(len(rows), 3, Mm(M + 137), Mm(y + 16), Mm(W - 137), Mm(54))
tbl = gf.table
tbl.first_row = tbl.horz_banding = False
for r, row in enumerate(rows):
    for c, val in enumerate(row):
        cell = tbl.cell(r, c)
        cell.text = val
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        f = p.runs[0].font
        f.name, f.size, f.bold = FONT, Pt(16), r == 0
        f.color.rgb = BLACK
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE
        cell_border(cell)

text(M + 137, y + 74, W - 137, 25, [
    ("Ergebnis:", {"bold": True, "after": 0}),
    "F · s ist bei allen Rampen ungefähr 9,8 N·m. "
    "Damit stimmt die Goldene Regel.",
])

# ---------- Alltag ----------
y = 324
heading(M, y, W, "Beispiele aus dem Alltag")
for i, cap in enumerate([
        "Passstraße: Die Kurven machen den Weg länger, dafür braucht man weniger Kraft.",
        "Rampe: Man schiebt länger, aber leichter."]):
    bx = M + i * (COL + 10)
    ph = rect(bx, y + 14, COL, 44, line=BLACK, dash=True, lw=1)
    tf = ph.text_frame
    tf.text = "Foto"
    run = tf.paragraphs[0].runs[0]
    run.font.name, run.font.size, run.font.color.rgb = FONT, Pt(16), BLACK
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    text(bx, y + 60, COL, 14, [cap], size=15)

# ---------- Quellen ----------
text(M, 405, W, 10, [
    "Quellen: Universum Physik (Schulbuch), S. 90–93; Fotos: eigene"], size=12)

prs.save("/home/user/Project/poster/Poster_Geneigte_Ebene.pptx")
print("ok")
